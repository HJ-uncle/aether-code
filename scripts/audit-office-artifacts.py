#!/usr/bin/env python3
"""Independent, read-only office artifact audit; never creates campaign deliverables.

Example: python scripts/audit-office-artifacts.py --root .ae/office-run \
    --output .e2e-tmp/office-format-report.json --render-pdf \
    --evidence-dir .e2e-tmp/office-render-evidence

Exit 1 means a failed check or incomplete scan, 2 means invalid arguments. Unsupported
calculation and unperformed visual review remain explicit even when exit code is 0.
Each parser runs serially in a hidden, time-bounded child to protect the campaign.
"""
from __future__ import annotations

import argparse
import datetime as dt
from decimal import Decimal
import hashlib
import importlib.util
import io
import json
import math
import os
from pathlib import Path, PurePosixPath
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.parse
import xml.etree.ElementTree as ET
import zipfile

FORMATS = {".docx", ".xlsx", ".pptx", ".pdf"}
MAX_FILE = 32 * 1024 * 1024
MAX_EXPANDED = 128 * 1024 * 1024
MAX_ENTRIES = 4096
MAX_CELLS = 100000
MAX_CONTENT_CELLS = 10000
MAX_TEXT = 120000
MAX_FILES = 2000
NS = {"r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
CREATION_FLAGS = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0


def check(result, name, status, **details):
    result["checks"].append({"id": name, "status": status, **details})


def scalar(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def digest(data):
    return "sha256:" + hashlib.sha256(data).hexdigest()


def file_digest(path):
    if path.stat().st_size > MAX_FILE:
        raise ValueError("file exceeds 32 MiB audit limit")
    with path.open("rb") as source:
        data = source.read(MAX_FILE + 1)
    if len(data) > MAX_FILE:
        raise ValueError("file grew beyond 32 MiB audit limit")
    return digest(data)


class TextBudget:
    def __init__(self):
        self.left = MAX_TEXT
        self.truncated = False

    def take(self, text):
        text = str(text or "")
        if len(text) > self.left:
            self.truncated = True
        kept = text[:self.left]
        self.left -= len(kept)
        return kept


def xml(data):
    # OOXML has no need for DTDs; rejecting them avoids entity expansion entirely.
    if re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", data, re.I):
        raise ValueError("DTD/entity declarations are not permitted in OOXML")
    return ET.fromstring(data)


def rel_target(rels_name, target):
    decoded = urllib.parse.unquote(target.split("#", 1)[0])
    if "\\" in decoded or urllib.parse.urlsplit(decoded).scheme:
        raise ValueError("invalid internal relationship target")
    if rels_name == "_rels/.rels":
        source_dir = ""
    else:
        source_dir = posixpath.dirname(posixpath.dirname(rels_name))
    normalized = posixpath.normpath(posixpath.join(source_dir, decoded))
    if decoded.startswith("/"):
        normalized = posixpath.normpath(decoded).lstrip("/")
    if normalized == ".." or normalized.startswith("../"):
        raise ValueError("relationship target escapes package")
    return normalized


def package_preflight(data, kind, result):
    if not data.startswith(b"PK\x03\x04"):
        raise ValueError("extension does not match a ZIP/OOXML container")
    package = zipfile.ZipFile(io.BytesIO(data))
    entries = package.infolist()
    names = [entry.filename for entry in entries]
    if len(entries) > MAX_ENTRIES or len(set(names)) != len(names):
        raise ValueError("too many ZIP entries or duplicate part names")
    expanded = sum(entry.file_size for entry in entries)
    if expanded > MAX_EXPANDED:
        raise ValueError("ZIP expanded size exceeds 128 MiB audit limit")
    for entry in entries:
        part = PurePosixPath(entry.filename)
        if (part.is_absolute() or ".." in part.parts or "\\" in entry.filename
                or ":" in entry.filename or entry.flag_bits & 1):
            raise ValueError("unsafe or encrypted ZIP entry: " + entry.filename)
        if entry.file_size > MAX_FILE:
            raise ValueError("individual package part exceeds 32 MiB")
    required = {"docx": "word/document.xml", "xlsx": "xl/workbook.xml",
                "pptx": "ppt/presentation.xml"}[kind]
    for name in ("[Content_Types].xml", "_rels/.rels", required):
        if name not in names:
            raise ValueError("missing required OOXML part: " + name)
    broken = package.testzip()
    if broken:
        raise ValueError("ZIP CRC failure: " + broken)
    external, missing = [], []
    trees = {}
    for name in names:
        if name.endswith((".xml", ".rels")):
            trees[name] = xml(package.read(name))
    types = trees["[Content_Types].xml"]
    content_types = {item.get("PartName", "").lstrip("/"): item.get("ContentType")
                     for item in types if item.tag.endswith("}Override")}
    expected_type = {"docx": "wordprocessingml.document.main+xml",
                     "xlsx": "spreadsheetml.sheet.main+xml",
                     "pptx": "presentationml.presentation.main+xml"}[kind]
    if not str(content_types.get(required, "")).endswith(expected_type):
        raise ValueError("OOXML main content type does not match extension")
    for name, tree in trees.items():
        if not name.endswith(".rels"):
            continue
        ids = set()
        for relation in tree:
            rid = relation.get("Id")
            if not rid or rid in ids:
                raise ValueError("missing or duplicate relationship ID in " + name)
            ids.add(rid)
            target = relation.get("Target", "")
            if relation.get("TargetMode") == "External":
                external.append({"part": name, "type": relation.get("Type"),
                                 "target": target[:1024]})
                continue
            if not target:
                missing.append({"part": name, "target": target})
                continue
            resolved = rel_target(name, target)
            if resolved not in names:
                missing.append({"part": name, "target": resolved})
    result["metadata"]["package"] = {
        "entryCount": len(entries), "uncompressedBytes": expanded,
        "externalRelationships": external[:100],
        "externalRelationshipCount": len(external),
        "embeddedParts": [name for name in names if "/embeddings/" in name][:100],
        "macroParts": [name for name in names if "vbaproject" in name.lower()]}
    check(result, "ooxml.container", "passed", entryCount=len(entries))
    check(result, "ooxml.relationships", "failed" if missing else "passed",
          missing=missing[:100])
    if missing:
        raise ValueError("OOXML contains unresolved internal relationships")
    if external:
        result["limitations"].append("External relationships recorded but never fetched.")
    return package, trees


def inspect_docx(data, result):
    from docx import Document
    doc = Document(io.BytesIO(data))
    budget = TextBudget()
    paragraphs = [{"index": i, "text": budget.take(p.text),
                   "style": p.style.name if p.style else None,
                   "alignment": scalar(p.alignment)}
                  for i, p in enumerate(doc.paragraphs[:10000])]
    tables = []
    for i, table in enumerate(doc.tables[:200]):
        tables.append({"index": i, "rowCount": len(table.rows),
                       "columnCount": len(table.columns),
                       "rows": [[budget.take(cell.text) for cell in row.cells[:100]]
                                for row in table.rows[:1000]]})
        budget.truncated |= len(table.rows) > 1000 or len(table.columns) > 100
    headers, footers, sections = [], [], []
    for section in doc.sections:
        headers.append([budget.take(p.text) for p in section.header.paragraphs])
        footers.append([budget.take(p.text) for p in section.footer.paragraphs])
        sections.append({key: round(getattr(section, key).pt, 2)
                         if getattr(section, key) is not None else None
                         for key in ("page_width", "page_height", "top_margin",
                                     "bottom_margin", "left_margin", "right_margin")})
    all_text = [p["text"] for p in paragraphs]
    all_text += [cell for table in tables for row in table["rows"] for cell in row]
    result["content"] = {"text": "\n".join(all_text), "paragraphs": paragraphs,
                         "tables": tables, "headers": headers, "footers": footers,
                         "truncated": budget.truncated or len(doc.paragraphs) > 10000
                         or len(doc.tables) > 200}
    result["metadata"].update({"paragraphCount": len(doc.paragraphs),
                               "tableCount": len(doc.tables), "sections": sections,
                               "inlineShapeCount": len(doc.inline_shapes)})
    check(result, "docx.parse", "passed")
    result["limitations"].extend([
        "DOCX pagination, overflow and typography have not been rendered or visually reviewed.",
        "DOCX extraction covers body paragraphs, top-level tables and default headers/footers; text boxes, nested tables, footnotes and other variants are not a full-fidelity extraction."])


def formula_expected(expression, cells, stack=None):
    """Deliberately small independent oracle, never Excel-wide recalculation."""
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    stack = set() if stack is None else stack
    match = re.fullmatch(r"(SUM|AVERAGE|MIN|MAX|COUNT)\(([^()]*)\)", expression.lstrip("="), re.I)
    if not match:
        raise NotImplementedError("only simple same-sheet SUM/AVERAGE/MIN/MAX/COUNT supported")
    numbers = []
    for argument in match.group(2).split(","):
        argument = argument.strip().replace("$", "")
        if not argument:
            raise NotImplementedError("empty function argument")
        if re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", argument):
            numbers.append(Decimal(argument))
            continue
        if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,6}(?::[A-Z]{1,3}[1-9]\d{0,6})?", argument, re.I):
            raise NotImplementedError("unsupported reference or expression: " + argument)
        lo_col, lo_row, hi_col, hi_row = range_boundaries(argument.upper())
        if (hi_col - lo_col + 1) * (hi_row - lo_row + 1) > MAX_CELLS:
            raise NotImplementedError("reference exceeds independent calculation bound")
        for row in range(lo_row, hi_row + 1):
            for col in range(lo_col, hi_col + 1):
                address = f"{get_column_letter(col)}{row}"
                cell = cells.get(address)
                if not cell:
                    continue
                value = cell["value"]
                if cell["formula"]:
                    if address in stack:
                        raise NotImplementedError("circular formula dependency")
                    if len(stack) > 64:
                        raise NotImplementedError("formula dependency depth limit")
                    value = formula_expected(cell["formula"], cells, stack | {address})
                elif cell["dataType"] == "d":
                    raise NotImplementedError("date/time reference requires Excel serial-date semantics")
                elif cell["dataType"] == "e":
                    raise NotImplementedError("referenced error cell needs Excel semantics")
                if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
                    numbers.append(Decimal(str(value)))
    fn = match.group(1).upper()
    if fn == "COUNT":
        return Decimal(len(numbers))
    if fn == "SUM":
        return sum(numbers, Decimal(0))
    if fn == "AVERAGE":
        if not numbers:
            raise NotImplementedError("AVERAGE of empty range has an error result")
        return sum(numbers, Decimal(0)) / len(numbers)
    return (min(numbers) if fn == "MIN" else max(numbers)) if numbers else Decimal(0)


def inspect_xlsx(data, trees, result):
    from openpyxl import load_workbook
    workbook_xml = trees["xl/workbook.xml"]
    relationships = trees.get("xl/_rels/workbook.xml.rels")
    if relationships is None:
        raise ValueError("workbook has no relationships")
    parts = {r.get("Id"): rel_target("xl/_rels/workbook.xml.rels", r.get("Target", ""))
             for r in relationships if r.get("TargetMode") != "External"}
    sheet_cells, total_cells = {}, 0
    for entry in workbook_xml.findall("s:sheets/s:sheet", NS):
        part = parts.get(entry.get("{" + NS["r"] + "}id"))
        if part not in trees:
            raise ValueError("worksheet part is missing")
        nodes = trees[part].findall("s:sheetData/s:row/s:c", NS)
        addresses = [node.get("r", "") for node in nodes]
        if len(set(addresses)) != len(addresses) or any(not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,6}", a) for a in addresses):
            raise ValueError("duplicate or invalid worksheet cell address")
        sheet_cells[entry.get("name")] = addresses
        total_cells += len(addresses)
    if total_cells > MAX_CELLS:
        raise ValueError("workbook exceeds 100000 populated-cell audit limit")
    wb = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    budget, sheets, layouts, formulas, remaining = TextBudget(), [], [], [], MAX_CONTENT_CELLS
    for sheet in wb.worksheets:
        cells = {}
        for address in sheet_cells.get(sheet.title, []):
            cell, cache = sheet[address], cached[sheet.title][address]
            raw_formula = cell.value if cell.data_type == "f" else None
            cells[address] = {"value": scalar(cell.value), "cachedValue": scalar(cache.value),
                              "formula": scalar(raw_formula), "dataType": cell.data_type,
                              "numberFormat": cell.number_format}
        for address, cell in cells.items():
            if cell["formula"] is None:
                continue
            item = {"sheet": sheet.title, "cell": address, "expression": cell["formula"],
                    "cache": cell["cachedValue"], "expected": None}
            try:
                expected = formula_expected(cell["formula"], cells, {address})
                item["expected"] = scalar(expected)
                if cell["cachedValue"] is None:
                    item.update(status="not-performed", reason="no saved formula cache")
                elif isinstance(cell["cachedValue"], (int, float)) and not isinstance(cell["cachedValue"], bool):
                    equal = abs(Decimal(str(cell["cachedValue"])) - expected) <= max(Decimal("1e-9"), abs(expected) * Decimal("1e-12"))
                    item.update(status="passed" if equal else "failed", reason="independent arithmetic versus saved cache")
                else:
                    item.update(status="failed", reason="saved cache is not the expected numeric type")
            except NotImplementedError as exc:
                item.update(status="unsupported", reason=str(exc))
            formulas.append(item)
            check(result, "xlsx.formula.cache", item["status"], **{k: v for k, v in item.items() if k != "status"})
        kept = {}
        for address, cell in list(cells.items())[:remaining]:
            kept[address] = {key: budget.take(value) if isinstance(value, str) else value
                             for key, value in cell.items()}
        remaining -= len(kept)
        sheets.append({"name": sheet.title, "cells": kept, "cellCount": len(cells),
                       "truncated": len(kept) != len(cells)})
        layouts.append({"name": sheet.title, "state": sheet.sheet_state,
                        "maxRow": sheet.max_row, "maxColumn": sheet.max_column,
                        "freezePanes": scalar(sheet.freeze_panes),
                        "autoFilter": sheet.auto_filter.ref, "printArea": scalar(sheet.print_area),
                        "mergedRanges": [str(r) for r in sheet.merged_cells.ranges][:1000],
                        "tableNames": list(sheet.tables)[:100],
                        "dataValidationCount": len(sheet.data_validations.dataValidation),
                        "conditionalFormatCount": len(sheet.conditional_formatting)})
    result["content"] = {"sheets": sheets, "truncated": total_cells > MAX_CONTENT_CELLS or budget.truncated,
                         "cellLimit": MAX_CONTENT_CELLS}
    result["metadata"].update({"sheetCount": len(wb.sheetnames), "cellCount": total_cells,
                               "sheets": layouts, "calculation": str(wb.calculation)})
    result["formulas"] = formulas
    check(result, "xlsx.parse", "passed", populatedCells=total_cells)
    result["limitations"].extend([
        "Formula checks cover only the listed simple same-sheet numeric functions; no Excel-wide recalculation was performed.",
        "Spreadsheet layout and print pagination have not been rendered or visually reviewed."])
    wb.close()
    cached.close()


def inspect_pptx(data, result):
    from pptx import Presentation
    presentation = Presentation(io.BytesIO(data))
    if len(presentation.slides) > 500:
        raise ValueError("presentation exceeds 500-slide audit bound")
    budget, slides, overflow = TextBudget(), [], []
    width, height = presentation.slide_width, presentation.slide_height
    for index, slide in enumerate(presentation.slides):
        texts, tables, shapes = [], [], []
        if len(slide.shapes) > 5000:
            raise ValueError("slide exceeds 5000-shape audit bound")
        for shape in slide.shapes:
            geometry = {"name": shape.name, "type": str(shape.shape_type),
                        "leftPt": round(shape.left.pt, 2), "topPt": round(shape.top.pt, 2),
                        "widthPt": round(shape.width.pt, 2), "heightPt": round(shape.height.pt, 2)}
            shapes.append(geometry)
            if shape.left < -12700 or shape.top < -12700 or shape.left + shape.width > width + 12700 or shape.top + shape.height > height + 12700:
                overflow.append({"slide": index + 1, **geometry})
            if shape.has_text_frame:
                texts.append(budget.take(shape.text))
            if shape.has_table:
                tables.append([[budget.take(cell.text) for cell in row.cells] for row in shape.table.rows])
        notes = budget.take(slide.notes_slide.notes_text_frame.text) if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ""
        slides.append({"number": index + 1, "text": "\n".join(texts), "tables": tables,
                       "shapes": shapes, "notes": notes})
    result["content"] = {"slides": slides, "truncated": budget.truncated}
    result["metadata"].update({"slideCount": len(slides), "widthPt": round(width.pt, 2),
                               "heightPt": round(height.pt, 2), "offSlideShapes": overflow[:100]})
    check(result, "pptx.parse", "passed", slideCount=len(slides))
    if overflow:
        result["limitations"].append("Some shapes extend beyond slide bounds; intent and clipping require visual review.")
    result["limitations"].extend([
        "PPTX typography, image fidelity and overflow have not been rendered or visually reviewed.",
        "PPTX extraction covers top-level text frames and tables; grouped text, chart labels and SmartArt may require separate inspection."])


def pdf_fonts(page):
    records = []
    resources = page.get("/Resources")
    resources = resources.get_object() if resources else {}
    fonts = resources.get("/Font", {})
    fonts = fonts.get_object() if hasattr(fonts, "get_object") else fonts
    for name, reference in fonts.items():
        font = reference.get_object()
        descriptor = font.get("/FontDescriptor")
        descendants = font.get("/DescendantFonts")
        if not descriptor and descendants:
            descriptor = descendants[0].get_object().get("/FontDescriptor")
        descriptor = descriptor.get_object() if descriptor else {}
        records.append({"resource": str(name), "baseFont": str(font.get("/BaseFont", "")),
                        "subtype": str(font.get("/Subtype", "")),
                        "embedded": any(key in descriptor for key in ("/FontFile", "/FontFile2", "/FontFile3")),
                        "hasToUnicode": "/ToUnicode" in font})
    return records


def inspect_pdf(data, result):
    from pypdf import PdfReader
    if not data.startswith(b"%PDF-"):
        raise ValueError("extension does not match PDF magic")
    reader = PdfReader(io.BytesIO(data), strict=True)
    if reader.is_encrypted:
        raise ValueError("encrypted PDF cannot be independently inspected without credentials")
    if len(reader.pages) > 500:
        raise ValueError("PDF exceeds 500-page audit limit")
    budget, pages = TextBudget(), []
    for index, page in enumerate(reader.pages):
        pages.append({"number": index + 1, "widthPt": float(page.mediabox.width),
                      "heightPt": float(page.mediabox.height),
                      "text": budget.take(page.extract_text() or ""), "fonts": pdf_fonts(page)})
    result["content"] = {"pages": pages, "text": "\n".join(p["text"] for p in pages),
                         "truncated": budget.truncated}
    result["metadata"].update({"pageCount": len(pages), "encrypted": False,
                               "documentInfo": {str(k): scalar(v) for k, v in (reader.metadata or {}).items()}})
    check(result, "pdf.parse", "passed", pageCount=len(pages))
    result["limitations"].append("Extracted PDF text does not prove Chinese glyph fidelity or table layout; raster images still require visual review.")


def capabilities():
    dependencies = {name: importlib.util.find_spec(name) is not None
                    for name in ("docx", "openpyxl", "pptx", "pypdf", "PIL")}
    poppler = shutil.which("pdftoppm")
    fallback = Path(sys.executable).resolve().parent.parent / "native/poppler/Library/bin/pdftoppm.exe"
    if not poppler and fallback.is_file():
        poppler = str(fallback)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    return {"python": sys.executable, "libraries": dependencies,
            "pdfRasterizer": {"path": poppler, "available": bool(poppler), "executed": False},
            "officeRenderer": {"path": soffice, "discoveredOnPath": bool(soffice),
                               "executed": False, "status": "not-performed"},
            "artifactTool": {"status": "not-performed", "reason": "not imported or exercised by this validator"}}


def render_pdf(data, result, args, caps):
    executable = caps["pdfRasterizer"]["path"]
    if not executable:
        check(result, "pdf.render", "not-performed", reason="pdftoppm not discovered")
        return
    from PIL import Image
    folder = Path(args.evidence_dir).resolve() / hashlib.sha256(data).hexdigest()
    folder.mkdir(parents=True, exist_ok=True)
    pages = min(args.max_render_pages, result["metadata"].get("pageCount", 0))
    if pages == 0:
        check(result, "pdf.render", "failed", reason="PDF has no pages")
        return
    # Render the already audited bytes so a concurrent source revision cannot mix evidence.
    with tempfile.TemporaryDirectory(prefix="pdf-audit-") as temporary:
        snapshot = Path(temporary) / "source.pdf"
        snapshot.write_bytes(data)
        command = [executable, "-f", "1", "-l", str(pages), "-r", "96", "-scale-to", "1600",
                   "-png", str(snapshot), str(folder / "page")]
        completed = subprocess.run(command, capture_output=True, timeout=35,
                                   creationflags=CREATION_FLAGS, check=False)
    if completed.returncode:
        check(result, "pdf.render", "failed", reason=completed.stderr.decode("utf-8", errors="replace")[:2000])
        return
    images = []
    for image_path in sorted(folder.glob("page-*.png"))[:pages]:
        with Image.open(image_path) as picture:
            picture.load()
            images.append({"path": str(image_path), "width": picture.width, "height": picture.height,
                           "sha256": digest(image_path.read_bytes())})
    status = "passed" if len(images) == pages else "failed"
    check(result, "pdf.render", status, requestedPages=pages, renderedPages=len(images))
    result["visual"] = {"status": "rendered" if status == "passed" else "failed",
                         "reviewStatus": "not-performed", "renderer": "pdftoppm",
                         "images": images, "totalPages": result["metadata"]["pageCount"],
                         "sampling": pages < result["metadata"]["pageCount"]}


def inspect_file(path, args, caps):
    result = {"path": str(path), "relativePath": path.relative_to(Path(args.root).resolve()).as_posix(),
              "format": path.suffix.lower().lstrip("."), "sha256": None, "sizeBytes": None,
              "structure": {"status": "not-performed"}, "content": {}, "metadata": {},
              "formulas": [], "checks": [], "limitations": [], "sourceUnchanged": None,
              "visual": {"status": "not-performed", "reviewStatus": "not-performed"}}
    data = None
    try:
        size = path.stat().st_size
        result["sizeBytes"] = size
        if size > MAX_FILE:
            raise ValueError("file exceeds 32 MiB audit limit")
        with path.open("rb") as source:
            data = source.read(MAX_FILE + 1)
        if len(data) > MAX_FILE:
            raise ValueError("file grew beyond 32 MiB audit limit")
        result["sha256"] = digest(data)
        kind = result["format"]
        if kind == "pdf":
            inspect_pdf(data, result)
        else:
            package, trees = package_preflight(data, kind, result)
            try:
                {"docx": lambda: inspect_docx(data, result),
                 "xlsx": lambda: inspect_xlsx(data, trees, result),
                 "pptx": lambda: inspect_pptx(data, result)}[kind]()
            finally:
                package.close()
        result["structure"]["status"] = "passed"
    except ImportError as exc:
        check(result, "parser.dependency", "unsupported", reason=str(exc))
        result["structure"]["status"] = "unsupported"
    except Exception as exc:
        check(result, "artifact.parse", "failed", reason=f"{type(exc).__name__}: {exc}")
        result["structure"]["status"] = "failed"
    if data is not None and result["structure"]["status"] == "passed" and result["format"] == "pdf" and args.render_pdf:
        try:
            render_pdf(data, result, args, caps)
        except Exception as exc:
            check(result, "pdf.render", "failed", reason=f"{type(exc).__name__}: {exc}")
    if result["sha256"] is not None:
        try:
            result["sourceUnchanged"] = file_digest(path) == result["sha256"]
            check(result, "artifact.source-stable", "passed" if result["sourceUnchanged"] else "failed")
        except Exception as exc:
            result["sourceUnchanged"] = False
            check(result, "artifact.source-stable", "failed", reason=str(exc))
    if result["content"].get("truncated"):
        result["limitations"].append("Extracted report content is bounded and truncated; it is not a complete-file content assertion.")
    if result["visual"]["status"] == "not-performed":
        check(result, "artifact.visual-review", "not-performed", reason="no visual inspection performed")
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", required=True)
    parser.add_argument("--output")
    parser.add_argument("--path", action="append", default=[])
    parser.add_argument("--render-pdf", action="store_true")
    parser.add_argument("--evidence-dir")
    parser.add_argument("--max-render-pages", type=int, default=3)
    parser.add_argument("--file-timeout", type=int, default=90)
    parser.add_argument("--worker", help=argparse.SUPPRESS)
    args = parser.parse_args()
    root = Path(args.root).resolve()
    if not root.is_dir():
        parser.error("--root must be an existing directory")
    if not 1 <= args.max_render_pages <= 10 or not 10 <= args.file_timeout <= 180:
        parser.error("render pages must be 1..10 and file timeout must be 10..180 seconds")
    if args.render_pdf and not args.evidence_dir:
        parser.error("--render-pdf requires --evidence-dir")
    if args.render_pdf and args.file_timeout < 45:
        parser.error("PDF rendering requires --file-timeout >= 45 seconds")
    caps = capabilities()
    if args.worker:
        candidate = Path(args.worker).resolve()
        if not candidate.is_relative_to(root) or candidate.suffix.lower() not in FORMATS:
            parser.error("worker source must be an office artifact beneath root")
        print(json.dumps(inspect_file(candidate, args, caps), ensure_ascii=True, allow_nan=False))
        return 0
    if not args.output:
        parser.error("--output is required")
    output = Path(args.output).resolve()
    if output.suffix.lower() != ".json":
        parser.error("--output must end in .json")
    candidates, scan_issues = [], []
    if args.path:
        for supplied in args.path:
            candidate = (root / supplied).resolve()
            if not candidate.is_relative_to(root) or candidate.suffix.lower() not in FORMATS:
                parser.error("--path must name an office artifact beneath root")
            candidates.append(candidate)
    else:
        for folder, dirs, names in os.walk(root, followlinks=False, onerror=lambda e: scan_issues.append(str(e))):
            dirs[:] = sorted(d for d in dirs if not (Path(folder) / d).is_symlink())
            for name in sorted(names):
                candidate = Path(folder) / name
                if candidate.suffix.lower() not in FORMATS:
                    continue
                if candidate.is_symlink() or not candidate.resolve().is_relative_to(root):
                    scan_issues.append("Skipped linked/out-of-root source: " + str(candidate))
                    continue
                candidates.append(candidate.resolve())
                if len(candidates) >= MAX_FILES:
                    scan_issues.append("Scan stopped at 2000-artifact safety bound; completeness not verified")
                    break
            if len(candidates) >= MAX_FILES:
                break
    files = []
    for path in sorted(set(candidates), key=lambda p: str(p).lower()):
        command = [sys.executable, str(Path(__file__).resolve()), "--root", str(root), "--worker", str(path),
                   "--file-timeout", str(args.file_timeout), "--max-render-pages", str(args.max_render_pages)]
        if args.render_pdf:
            command.extend(["--render-pdf", "--evidence-dir", str(Path(args.evidence_dir).resolve())])
        try:
            completed = subprocess.run(command, capture_output=True, timeout=args.file_timeout,
                                       creationflags=CREATION_FLAGS, check=False)
            if completed.returncode:
                raise RuntimeError(completed.stderr.decode("utf-8", errors="replace")[:2000])
            files.append(json.loads(completed.stdout))
        except Exception as exc:
            files.append({"path": str(path), "relativePath": path.relative_to(root).as_posix(),
                          "format": path.suffix[1:].lower(), "sha256": None, "sizeBytes": None,
                          "structure": {"status": "failed"}, "sourceUnchanged": None,
                          "content": {}, "metadata": {}, "formulas": [], "limitations": [],
                          "visual": {"status": "not-performed", "reviewStatus": "not-performed"},
                          "checks": [{"id": "parser.worker", "status": "failed", "reason": str(exc)}]})
    statuses = {status: sum(c["status"] == status for f in files for c in f["checks"])
                for status in ("passed", "failed", "not-performed", "unsupported")}
    if not files:
        scan_issues.append("No office artifacts found; empty scan is not successful acceptance")
    report = {"schemaVersion": 1, "generatedAt": dt.datetime.now(dt.timezone.utc).isoformat(),
              "root": str(root), "capabilities": caps, "files": files,
              "summary": {"fileCount": len(files), "checks": statuses, "scanIssues": scan_issues,
                          "failedFiles": sum(any(c["status"] == "failed" for c in f["checks"]) for f in files),
                          "structurallyPassedFiles": sum(f["structure"]["status"] == "passed" for f in files),
                          "visualReviewedFiles": 0, "businessAcceptance": "not-performed"}}
    caps["pdfRasterizer"]["executed"] = any(c["id"] == "pdf.render" and c["status"] in ("passed", "failed") for f in files for c in f["checks"])
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(json.dumps({"report": str(output), **report["summary"]}, ensure_ascii=True))
    return 1 if statuses["failed"] or scan_issues else 0


if __name__ == "__main__":
    raise SystemExit(main())
